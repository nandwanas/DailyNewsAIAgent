import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config.settings import Config
from src.utils import today_str

logger = logging.getLogger("excel_exporter")

HEADER_FILL = PatternFill("solid", fgColor="1B5E20")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
POS_FILL = PatternFill("solid", fgColor="E8F5E9")
NEG_FILL = PatternFill("solid", fgColor="FFEBEE")
NEU_FILL = PatternFill("solid", fgColor="FFF8E1")


def _style_sheet(ws):
    """Apply header styling and auto-fit columns."""
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Auto-fit columns (cap at 60 chars wide)
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def _build_raw_df(articles: list) -> pd.DataFrame:
    rows = []
    for a in articles:
        rows.append({
            "article_id": a.article_id,
            "title": a.title,
            "url": a.url,
            "source_name": a.source_name,
            "published_at": a.published_at,
            "fetch_source": a.fetch_source,
            "image_url": a.image_url or "",
            "category": a.category,
        })
    return pd.DataFrame(rows)


def _build_analysis_df(articles: list) -> pd.DataFrame:
    rows = []
    for a in articles:
        kp = a.key_points or []
        tg = a.tags or []
        rows.append({
            "title": a.title,
            "subcategory": a.subcategory,
            "summary_en": a.summary_en,
            "key_point_1": kp[0] if len(kp) > 0 else "",
            "key_point_2": kp[1] if len(kp) > 1 else "",
            "key_point_3": kp[2] if len(kp) > 2 else "",
            "sentiment": a.sentiment,
            "sentiment_score": round(a.sentiment_score, 2),
            "importance_score": a.importance_score,
            "tag_1": tg[0] if len(tg) > 0 else "",
            "tag_2": tg[1] if len(tg) > 1 else "",
            "tag_3": tg[2] if len(tg) > 2 else "",
            "processing_errors": "; ".join(a.processing_errors) if a.processing_errors else "",
        })
    return pd.DataFrame(rows)


def _build_translations_df(articles: list, lang_names: dict) -> pd.DataFrame:
    rows = []
    for a in articles:
        for lang_code, t in (a.translations or {}).items():
            rows.append({
                "article_id": a.article_id,
                "title_en": a.title,
                "lang_code": lang_code,
                "language_name": lang_names.get(lang_code, lang_code),
                "title_translated": t.get("title", ""),
                "summary_translated": t.get("summary", ""),
                "key_points_translated": " | ".join(t.get("key_points", [])),
            })
    return pd.DataFrame(rows) if rows else pd.DataFrame(columns=[
        "article_id", "title_en", "lang_code", "language_name",
        "title_translated", "summary_translated", "key_points_translated"
    ])


class ExcelExporter:
    def __init__(self, config: Config):
        self.config = config
        self.output_dir = config.output_dir / "excel"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(self, articles: list) -> Path:
        run_date = today_str()
        out_path = self.output_dir / f"{run_date}_agriculture_data.xlsx"

        raw_df = _build_raw_df(articles)
        analysis_df = _build_analysis_df(articles)
        translations_df = _build_translations_df(articles, self.config.LANGUAGE_NAMES)

        with pd.ExcelWriter(str(out_path), engine="openpyxl") as writer:
            raw_df.to_excel(writer, sheet_name="RawData", index=False)
            analysis_df.to_excel(writer, sheet_name="AIAnalysis", index=False)
            translations_df.to_excel(writer, sheet_name="Translations", index=False)

        # Apply formatting with openpyxl
        wb = load_workbook(str(out_path))
        for sheet_name in wb.sheetnames:
            _style_sheet(wb[sheet_name])

        # Colour-code sentiment rows in AIAnalysis
        ws_analysis = wb["AIAnalysis"]
        sent_col = None
        for cell in ws_analysis[1]:
            if cell.value == "sentiment":
                sent_col = cell.column
                break
        if sent_col:
            for row in ws_analysis.iter_rows(min_row=2):
                val = row[sent_col - 1].value or ""
                fill = POS_FILL if val == "positive" else (NEG_FILL if val == "negative" else NEU_FILL)
                for cell in row:
                    cell.fill = fill

        wb.save(str(out_path))
        logger.info(f"Excel saved: {out_path} ({len(articles)} articles)")
        return out_path
